#!/usr/bin/env python3
"""Validate package fidelity, reproducible outputs, examples and portable links."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from prepare_inputs import prepare, read_csv
from build_collaborator_report import opposition_tier, selection_key
from validate_results import validate
from ethics_analysis.schema import PARTITIONS


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--results", type=Path, default=ROOT / "results")
    ap.add_argument("--source", type=Path, help="Optional original repository for complete source-value comparison")
    ap.add_argument("--reference-checks", action="store_true")
    args = ap.parse_args()
    results = args.results.resolve()
    output = ROOT / "provenance/VALIDATION.json" if results == ROOT/"results" else results/"PACKAGE_VALIDATION.json"
    checks, source_comparison = 0, {}

    def check(condition, message):
        nonlocal checks
        if not condition:
            raise ValueError(message)
        checks += 1

    config = json.loads((ROOT / "config/model_groups.json").read_text())
    groups = config["groups"]
    check(len(groups["frontier"]) == 3 and len(groups["other_models"]) == 6,
          "Wrong frontier/remaining panel sizes")
    check(set(groups["frontier"]).isdisjoint(groups["other_models"]) and
          set(groups["frontier"]+groups["other_models"]) == set(groups["all"]), "Model panels do not partition all nine")
    schema = json.loads((ROOT / "docs/schema.json").read_text())
    human_lookup = {}
    for p in PARTITIONS:
        rows = read_csv(ROOT / "data/humans" / f"{p}.csv")
        check(len(rows) == 1995, f"Wrong human rows: {p}")
        for row in rows:
            check(row["item_id"] not in human_lookup, "Duplicate human item")
            human_lookup[row["item_id"]] = row
            check(all(row[f"annotator_{i}_source_cell_input"] for i in (1,2,3)), "Missing original input cell")
            check(all(row[f"annotator_{i}_source_cell_output"] for i in (1,2,3)), "Missing original output cell")
            check(len(json.loads(row["all_primary_rationale_sets_json"])) == 3,
                  "Missing one of the three primary rationale sets")
    with tempfile.TemporaryDirectory(prefix="ethics_package_check_") as tmp:
        human, llm = prepare(Path(tmp))
        rebuilt = read_csv(human / "data/processed/individual/all_annotations.csv")
        check(len(rebuilt) == 29925 and len(human_lookup) == 9975, "Wrong corpus totals")
        check(set(Counter(r["annotator_code"] for r in rebuilt).values()) == {1575}, "Wrong per-person workload")
        for filename in ("majority_labels.csv", "categorical_rationales.csv"):
            check(len(read_csv(human / "data/processed/aggregated" / filename)) == 9975, "Incomplete aggregate")
        if args.source:
            source = args.source.resolve()
            original = read_csv(source / "data/processed/individual/all_annotations.csv")
            lookup = {(r["partition"],r["task_instruction_id"],r["annotation_copy"]):r for r in original}
            for row in rebuilt:
                key = row["partition"],row["task_instruction_id"],row["annotation_copy"]
                check(row == lookup[key], f"Normalized source annotation differs: {key}")
            for entry in read_csv(ROOT / "provenance/SOURCE_MANIFEST.csv"):
                check(hashlib.sha256((source/entry["source_relative_path"]).read_bytes()).hexdigest() == entry["sha256"],
                      "Original source file changed")
            from openpyxl import load_workbook
            from export_from_source import HEADERS
            original_cells = {}
            import warnings
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
            for name in {r["source_workbook"] for r in rebuilt}:
                book = load_workbook(source / "data/raw/annotations" / name, read_only=True, data_only=False)
                it = book.worksheets[0].iter_rows(values_only=True)
                header = next(it)
                cols = {HEADERS[h]:i for i,h in enumerate(header) if h in HEADERS}
                for row_no, values in enumerate(it,2):
                    original_cells[name,str(row_no)] = {k:"" if values[i] is None else str(values[i]) for k,i in cols.items()}
                book.close()
            for item in human_lookup.values():
                for slot in (1,2,3):
                    cells = original_cells[item[f"annotator_{slot}_source_workbook"], item[f"annotator_{slot}_source_excel_row"]]
                    check(all(item[f"annotator_{slot}_source_cell_{field}"] == cells.get(field,"") for field in schema["source_cell_fields"]),
                          "Original source-cell values differ")
            source_comparison = dict(normalized_records_compared=29925, original_cell_records_compared=29925,
                                     original_source_hashes_unchanged=True)
            # Reconcile newly packaged model labels/rationale point estimates with the previous complete analysis.
            previous = source / "reports/production/analysis_results_20260922"
            for rel, keys, metrics in (("judgments/model_vs_human.csv",("model","sample","partition"),("n","matches","agreement","kappa")),
                                       ("rationales/model_human_overlap.csv",("model","field","scope"),("n_items","jaccard","dice","exact_set"))):
                saved = {tuple(r[k] for k in keys):r for r in read_csv(previous/rel)}
                for row in read_csv(results/rel):
                    old = saved[tuple(row[k] for k in keys)]
                    check(all(abs(float(row[m])-float(old[m])) < 1e-10 for m in metrics), "Packaged result differs from source analysis")
            source_comparison["previous_label_and_rationale_results_reconciled"] = True

    validate(results, args.reference_checks)
    for rel, digest in json.loads((results / "PACKAGE_INPUT_CHECKSUMS.json").read_text()).items():
        check(hashlib.sha256((ROOT/rel).read_bytes()).hexdigest() == digest, "Frozen package CSV changed")
    selected = read_csv(results / "examples/selected_cases.csv")
    candidates = read_csv(results / "examples/all_disagreement_candidates.csv")
    for rows in (selected,candidates):
        for row in rows:
            for key in ("n_valid_models","n_opposing_models","n_valid_frontier","n_opposing_frontier","human_ambiguity_votes","human_cultural_flag_votes","selection_tier"):
                row[key] = int(row[key])
            row["human_unanimous"] = row["human_unanimous"] == "True"
            item = human_lookup[row["item_id"]]
            expected = item["human_majority_label"]
            responses = json.loads(row["model_annotations_json"])
            good = [r for r in responses.values() if r and r["label_status"] == "valid"]
            check(len(good) == row["n_valid_models"] and sum(r["normalized_label"] != expected for r in good) == row["n_opposing_models"],
                  "Disagreement votes do not match evidence")
            check(opposition_tier(row) == row["selection_tier"] and row["human_unanimous"] and item["human_unanimous"] == "TRUE",
                  "Selected case fails unanimous-human opposition predicate")
            check(row["input"] == item["input"] and len(json.loads(row["human_annotations_json"])) == 3,
                  "Disagreement example lacks full original instruction/three annotations")
    for p in PARTITIONS:
        expected = sorted((r for r in candidates if r["partition"] == p),key=selection_key)[:config["examples_per_partition"]]
        actual = [r for r in selected if r["partition"] == p]
        check([r["item_id"] for r in expected] == [r["item_id"] for r in actual], "Disagreement selection is not deterministic")

    for path in (ROOT / "README.md", ROOT / "docs/DATA_DICTIONARY.md", ROOT / "docs/METHODS.md", results / "COLLABORATOR_REPORT.md", results / "examples/EXAMPLES.md"):
        content = path.read_text()
        check("/Users/" not in content, "Nonportable machine-specific path in a report")
        for target in re.findall(r"\]\(([^)]+)\)",content):
            if not target.startswith(("http://","https://","#")):
                resolved = (path.parent/target).resolve()
                check(resolved.exists() or resolved == output, f"Unresolved report link: {target}")
        width = None
        for line in content.splitlines():
            if not line.startswith("|"):
                width = None
            else:
                n = len(re.split(r"(?<!\\)\|",line)[1:-1])
                width = n if width is None else width
                check(n == width, "Malformed Markdown table")
    check(not any((ROOT/r).exists() for r in ("api.key","annotator_email_codes.csv","student_work.csv")),
          "Private administrative material present")
    check(not any(ROOT.rglob("*.xlsm")), "Original Office workbooks should not be distributed")
    result = dict(status="passed", package_checks=checks, source_comparison=source_comparison,
                  n_items=9975, human_ratings=29925, annotators=19, models=9, selected_examples=len(selected),
                  input_files=10, offline=True, source_files_modified=False,
                  git_initialized=(ROOT/".git").exists(),
                  reference_library_checks=args.reference_checks)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))


if __name__ == "__main__":
    main()
