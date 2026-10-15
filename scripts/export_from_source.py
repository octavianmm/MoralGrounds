#!/usr/bin/env python3
"""One-time, read-only export from the source repository into this package."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
import warnings
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "support" / "aggregation"))
from aggregate_majority_labels import SOURCE_FIELDS
from concatenate_categorical_rationales import categorical_structure
from ethics_analysis.ingest import load_data
from ethics_analysis.schema import PARTITIONS

RAW_FIELDS = ("input", "output", "rationale", "main_rationale", "moral_rationale", "ambiguity", "cultural_specificity")
HEADERS = {"Input": "input", "Output": "output", "Rationale": "rationale",
           "MAIN RATIONALES": "main_rationale", "MORAL RATIONALE": "moral_rationale",
           "Ambiguity": "ambiguity", "Cultural-specificity": "cultural_specificity"}


def compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise ValueError(f"Refusing to replace an existing package input: {path}")
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--source", type=Path, required=True)
    args = ap.parse_args()
    source = args.source.resolve()
    manifest = []

    def record(path):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        manifest.append(dict(source_relative_path=str(path.relative_to(source)), sha256=digest,
                             bytes=path.stat().st_size))

    input_path = source / "data/processed/individual/all_annotations.csv"
    with input_path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        original_fields = reader.fieldnames
        original = list(reader)
    record(input_path)
    items, humans, models, model_names, groups, audit = load_data(str(source), str(source / "ETHICS10k"))
    assert len(items) == 9975 and len(humans) == 29925 and len(models) == 89752
    assert len(model_names) == 9 and len(groups["frontier"]) == 3 and len(groups["other_models"]) == 6
    for entry in audit.source_manifest:
        base = source if entry["source"] == source.name else source / "ETHICS10k"
        record(base / entry["file"])
    grouped = defaultdict(list)
    for row in original:
        grouped[row["partition"], row["task_instruction_id"]].append(row)

    # Only cell values are read. Office authors, cached save paths and VBA are never exported.
    from openpyxl import load_workbook
    warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
    raw_cells = {}
    for name in sorted({r["source_workbook"] for r in original}):
        path = source / "data/raw/annotations" / name
        record(path)
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator)
        columns = {i: HEADERS[h] for i, h in enumerate(headers) if h in HEADERS}
        for excel_row, values in enumerate(iterator, 2):
            raw_cells[name, str(excel_row)] = {field: "" if values[i] is None else str(values[i]) for i, field in columns.items()}
        book.close()
    extras = [f for f in original_fields if f not in SOURCE_FIELDS]
    for p in PARTITIONS:
        output = []
        for (partition, tid), bundle in sorted(grouped.items(), key=lambda v: (v[0][0], int(v[0][1]))):
            if partition != p:
                continue
            bundle.sort(key=lambda r: int(r["annotation_copy"]))
            assert [int(r["annotation_copy"]) for r in bundle] == [1, 2, 3]
            assert len({r["annotator_code"] for r in bundle}) == 3
            first = bundle[0]
            assert all(len({r[f] for r in bundle}) == 1 for f in SOURCE_FIELDS)
            votes = Counter(r["main_label"] for r in bundle)
            majority, majority_n = votes.most_common(1)[0]
            all_rationales = categorical_structure(p, bundle)
            row = {f: first[f] for f in SOURCE_FIELDS}
            row.update(item_id=f"{p}:{tid}", human_majority_label=majority,
                       human_majority_count=majority_n, human_unanimous="TRUE" if majority_n == 3 else "FALSE",
                       individual_labels_json=compact([r["main_label"] for r in bundle]),
                       annotator_codes_json=compact([r["annotator_code"] for r in bundle]),
                       all_primary_rationale_sets_json=compact(all_rationales["categorical_rationales"]),
                       all_secondary_rationale_sets_json=compact(all_rationales["secondary_categorical_rationales"]),
                       all_free_text_rationales_json=compact([r["rationale"] for r in bundle] if p == "virtue" else []),
                       majority_aligned_copies_json=compact([int(r["annotation_copy"]) for r in bundle if r["main_label"] == majority]))
            for slot, annotation in enumerate(bundle, 1):
                for field in extras:
                    row[f"annotator_{slot}_{field}"] = annotation[field]
                cells = raw_cells[annotation["source_workbook"], annotation["source_excel_row"]]
                for field in RAW_FIELDS:
                    row[f"annotator_{slot}_source_cell_{field}"] = cells.get(field, "")
            output.append(row)
        assert len(output) == 1995
        write_csv(ROOT / "data/humans" / f"{p}.csv", output)

    for p in PARTITIONS:
        output = []
        for obs in sorted((o for o in models if o["partition"] == p), key=lambda o: (o["actor"], int(o["item_id"].split(":")[1]))):
            fields = obs["fields"]
            raw = lambda name: fields.get(name, {}).get("raw", "")
            primary = {"commonsense": "commonsense_principles", "justice": "justice_rationale", "utilitarianism": "utilitarianism_rationales", "virtue": "virtue_text"}.get(p)
            row = dict(item_id=obs["item_id"], partition=p, model=obs["actor"],
                       raw_output=obs["raw_label"], raw_rationale=raw(primary) if primary else "",
                       raw_main_rationale=raw("deontology_main"), raw_moral_rationale=raw("deontology_principles"),
                       normalized_label=({"commonsense": ("NOT_MORALLY_WRONG", "MORALLY_WRONG"), "utilitarianism": ("A", "B")}.get(p, ("NO", "YES"))[obs["label"]] if obs["label"] is not None else ""),
                       label_status=obs["label_status"], structured_fields_json=compact(fields),
                       source_export=obs["source_file"], source_row=obs["source_row"])
            output.append(row)
        write_csv(ROOT / "data/llms" / f"{p}.csv", output)
    config = dict(groups=groups, bootstrap=5000, seed=20260917, examples_per_partition=5,
                  interpretation="Fixed three commercial/frontier conditions versus the remaining six; not inferred regional populations.")
    (ROOT / "config/model_groups.json").write_text(json.dumps(config, indent=2)+"\n")
    (ROOT / "docs/schema.json").write_text(json.dumps(dict(schema_version=1, source_fields=list(SOURCE_FIELDS),
        original_individual_columns=original_fields, annotator_specific_fields=extras, source_cell_fields=list(RAW_FIELDS)), indent=2)+"\n")
    unique_manifest = {r["source_relative_path"]: r for r in manifest}
    for entry in unique_manifest.values():
        assert hashlib.sha256((source / entry["source_relative_path"]).read_bytes()).hexdigest() == entry["sha256"]
    target = ROOT / "provenance"
    target.mkdir(exist_ok=True)
    write_csv(target / "SOURCE_MANIFEST.csv", list(unique_manifest.values()))
    (target / "source_projection.json").write_text(json.dumps(dict(human_rows=29925, human_item_rows=9975,
        raw_workbooks_read=95, source_cells_retained=True, dissenting_annotations_retained=True,
        model_rows=89752, n_models=9, source_files_modified=False,
        privacy_exclusions=["student identity mapping", "emails", "payment records", "Office identity metadata", "API credentials"]), indent=2)+"\n")
    print("Exported five 1,995-row human files and five LLM files; all source hashes unchanged.")


if __name__ == "__main__":
    main()
